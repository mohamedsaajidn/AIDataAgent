import io

import openpyxl
import pandas as pd


def load_tabular(filename: str, content: bytes) -> dict[str, pd.DataFrame]:
    """Parse an uploaded CSV or XLSX file into one DataFrame per logical table.

    A CSV always yields a single sheet keyed by the file's stem. An XLSX file
    yields one entry per non-empty worksheet, since each sheet in a messy
    export usually represents a distinct table (e.g. "Customers", "Orders").
    """
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "csv":
        return {_stem(filename): _read_csv(content)}
    if ext == "xlsx":
        return _read_excel(content)
    raise ValueError(f"Unsupported file type: {filename} (only .csv and .xlsx are supported)")


def _read_csv(content: bytes) -> pd.DataFrame:
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            return pd.read_csv(io.BytesIO(content), encoding=encoding, low_memory=False)
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode CSV as utf-8 or latin-1")


def _read_excel(content: bytes) -> dict[str, pd.DataFrame]:
    workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    sheets: dict[str, pd.DataFrame] = {}
    for sheet_name in workbook.sheetnames:
        df = _sheet_to_dataframe(workbook[sheet_name])
        if not df.empty:
            sheets[sheet_name] = df
    return sheets


def _sheet_to_dataframe(worksheet) -> pd.DataFrame:
    _unmerge_cells(worksheet)
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        return pd.DataFrame()

    header_idx = _detect_header_row(rows)
    header = _dedupe_columns(
        [
            str(v).strip() if v is not None and str(v).strip() != "" else f"column_{i}"
            for i, v in enumerate(rows[header_idx])
        ]
    )
    data_rows = rows[header_idx + 1 :]
    df = pd.DataFrame(data_rows, columns=header)
    return df.dropna(how="all")


def _unmerge_cells(worksheet) -> None:
    # A merged header cell (e.g. a title spanning several columns) otherwise
    # leaves every cell but the top-left one blank once we read row values.
    for merged_range in list(worksheet.merged_cells.ranges):
        top_left_value = worksheet.cell(merged_range.min_row, merged_range.min_col).value
        worksheet.unmerge_cells(str(merged_range))
        for row in worksheet.iter_rows(
            min_row=merged_range.min_row,
            max_row=merged_range.max_row,
            min_col=merged_range.min_col,
            max_col=merged_range.max_col,
        ):
            for cell in row:
                cell.value = top_left_value


def _detect_header_row(rows: list[tuple], max_scan: int = 20) -> int:
    best_idx, best_score = 0, -1.0
    for idx, row in enumerate(rows[:max_scan]):
        score = _score_header_row(row)
        if score > best_score:
            best_score, best_idx = score, idx
    return best_idx


def _score_header_row(row: tuple) -> float:
    non_null = [v for v in row if v is not None and str(v).strip() != ""]
    if not non_null:
        return 0.0
    non_null_frac = len(non_null) / len(row)
    string_frac = sum(1 for v in non_null if isinstance(v, str)) / len(non_null)
    unique_frac = len(set(non_null)) / len(non_null)
    return non_null_frac * string_frac * unique_frac


def _dedupe_columns(names: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    result = []
    for name in names:
        if name in seen:
            seen[name] += 1
            result.append(f"{name}_{seen[name]}")
        else:
            seen[name] = 0
            result.append(name)
    return result


def _stem(filename: str) -> str:
    return filename.rsplit(".", 1)[0] if "." in filename else filename
